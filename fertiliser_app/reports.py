"""Excel exports."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2F6B3A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MONEY = '#,##0.00'
QTY = '#,##0.##'

KIND_LABELS = {"delivery": "Delivery", "usage": "Used", "adjustment": "Stock adjustment"}


def _header(ws, row, titles):
    for col, title in enumerate(titles, start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _movements_sheet(ws, movements, currency):
    cur = f" ({currency})" if currency else ""
    titles = [
        "Date / time", "Fertiliser", "Type", "Bags (+in / -out)", "Kg",
        f"Cost per bag{cur}", f"Value{cur}", "Recorded by", "Note",
        "Voided at", "Voided by", "Void reason",
    ]
    _header(ws, 1, titles)
    for r, m in enumerate(movements, start=2):
        values = [
            m["created_at"], m["fertiliser_name"], KIND_LABELS[m["kind"]], m["bags"],
            m["bags"] * m["kg_per_bag"], m["cost_per_bag"], m["bags"] * m["cost_per_bag"],
            m["recorded_by"], m["note"], m["voided_at"] or "", m["voided_by"] or "",
            m["void_reason"] or "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (4, 5):
                cell.number_format = QTY
            elif c in (6, 7):
                cell.number_format = MONEY
            if m["voided_at"]:
                cell.font = Font(strike=True, color="999999")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{max(1, len(movements) + 1)}"
    _widths(ws, [19, 24, 16, 12, 10, 12, 12, 16, 30, 19, 14, 24])


def monthly_report(summary, movements, title, currency):
    wb = Workbook()
    cur = f" ({currency})" if currency else ""
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Voided entries are excluded from totals. Closing stock is valued at the price in force at month end."
    ws["A2"].font = Font(italic=True, color="666666")

    titles = [
        "Fertiliser", "Kg per bag", "Opening stock (bags)", "Received (bags)",
        f"Received cost{cur}", "Used (bags)", "Used (kg)", f"Cost of fertiliser used{cur}",
        "Adjustments (bags)", "Closing stock (bags)", "Closing stock (kg)",
        f"Price per bag{cur}", f"Closing stock value{cur}",
    ]
    _header(ws, 4, titles)
    ws.row_dimensions[4].height = 45
    keys = [
        "name", "kg_per_bag", "opening_bags", "received_bags", "received_cost", "used_bags",
        "used_kg", "used_cost", "adjusted_bags", "closing_bags", "closing_kg",
        "price_per_bag", "closing_value",
    ]
    money_cols = {5, 8, 12, 13}
    first = 5
    for r, row in enumerate(summary, start=first):
        for c, key in enumerate(keys, start=1):
            cell = ws.cell(row=r, column=c, value=row[key])
            if c > 1:
                cell.number_format = MONEY if c in money_cols else QTY
    total_row = first + len(summary)
    ws.cell(row=total_row, column=1, value="TOTAL").font = BOLD
    for c in (5, 7, 8, 13):
        col = get_column_letter(c)
        cell = ws.cell(
            row=total_row, column=c,
            value=f"=SUM({col}{first}:{col}{total_row - 1})" if summary else 0,
        )
        cell.font = BOLD
        cell.number_format = MONEY if c in money_cols else QTY
    ws.freeze_panes = "B5"
    _widths(ws, [24, 9, 11, 11, 12, 10, 10, 14, 12, 11, 11, 11, 14])

    _movements_sheet(wb.create_sheet("All movements"), movements, currency)
    _movements_sheet(
        wb.create_sheet("Usage log"), [m for m in movements if m["kind"] == "usage"], currency
    )
    return _save(wb)


def movements_report(movements, currency):
    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    _movements_sheet(ws, movements, currency)
    return _save(wb)


def _save(wb):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
